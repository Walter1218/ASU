"""PPTX 缩略图生成器 — 从 .pptx 文件提取每页幻灯片的缩略图

实现方案：
- 使用 python-pptx 读取幻灯片结构信息，绘制简化缩略图
- 每页幻灯片生成一个 QPixmap（用于 PyQt6 缩略图条）
"""

import os
import logging
from typing import Optional

from PyQt6.QtGui import QPixmap, QPainter, QColor, QFont, QPen, QBrush
from PyQt6.QtCore import Qt, QRectF

logger = logging.getLogger(__name__)

# 缩略图尺寸
THUMB_WIDTH = 120
THUMB_HEIGHT = int(THUMB_WIDTH * 7.5 / 13.333)  # 16:9 比例

# 颜色方案
COLOR_BG = QColor(250, 250, 252)
COLOR_ACCENT = QColor(0, 82, 204)
COLOR_TITLE = QColor(15, 25, 45)
COLOR_TEXT = QColor(60, 64, 67)
COLOR_LIGHT_BG = QColor(235, 240, 248)


class PptxThumbnailGenerator:
    """从 .pptx 文件生成缩略图"""

    def __init__(self, thumb_width: int = THUMB_WIDTH, thumb_height: int = THUMB_HEIGHT):
        self.thumb_width = thumb_width
        self.thumb_height = thumb_height

    def generate_all(self, pptx_path: str) -> list:
        """为 .pptx 文件中的每一页生成缩略图

        Args:
            pptx_path: .pptx 文件路径

        Returns:
            list of QPixmap（每页一个）
        """
        if not os.path.isfile(pptx_path):
            logger.warning("PptxThumbnail: file not found: %s", pptx_path)
            return []

        try:
            from pptx import Presentation
            prs = Presentation(pptx_path)
            thumbnails = []

            for i, slide in enumerate(prs.slides):
                pixmap = self._render_slide_thumbnail(slide, i, len(prs.slides))
                thumbnails.append(pixmap)

            logger.info("PptxThumbnail: generated %d thumbnails from %s",
                        len(thumbnails), os.path.basename(pptx_path))
            return thumbnails

        except Exception as e:
            logger.error("PptxThumbnail: failed: %s", e)
            return []

    def generate_from_slides_data(self, slides_data: list) -> list:
        """从 slides_data 直接生成缩略图（无需 .pptx 文件）

        Args:
            slides_data: 幻灯片数据列表

        Returns:
            list of QPixmap
        """
        thumbnails = []
        for i, slide_data in enumerate(slides_data):
            pixmap = self._render_from_data(slide_data, i, len(slides_data))
            thumbnails.append(pixmap)
        return thumbnails

    # ── 内部渲染 ──────────────────────────────────────────────

    def _render_slide_thumbnail(self, slide, index: int, total: int) -> QPixmap:
        """渲染单张幻灯片的缩略图"""
        pixmap = QPixmap(self.thumb_width, self.thumb_height)
        pixmap.fill(COLOR_BG)

        painter = QPainter(pixmap)
        painter.setRenderHint(QPainter.RenderHint.Antialiasing)

        # 提取文本内容
        title = ""
        body_texts = []

        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if not text:
                    continue
                if shape == slide.shapes.title or not title:
                    title = text
                else:
                    body_texts.append(text)

        # 绘制底色
        painter.fillRect(0, 0, self.thumb_width, self.thumb_height, COLOR_BG)

        # 绘制标题区域
        if title:
            painter.fillRect(0, 0, self.thumb_width, self.thumb_height // 3, COLOR_LIGHT_BG)
            painter.setPen(QPen(COLOR_TITLE))
            font = QFont()
            font.setPointSize(7)
            font.setBold(True)
            painter.setFont(font)
            title_rect = QRectF(4, 2, self.thumb_width - 8, self.thumb_height // 3 - 4)
            painter.drawText(title_rect, Qt.TextFlag.TextWordWrap | Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter, title[:30])

        # 绘制正文
        if body_texts:
            painter.setPen(QPen(COLOR_TEXT))
            font = QFont()
            font.setPointSize(5)
            painter.setFont(font)
            y_offset = self.thumb_height // 3 + 4
            for text in body_texts[:4]:  # 最多显示 4 行
                text_rect = QRectF(6, y_offset, self.thumb_width - 12, 12)
                painter.drawText(text_rect, Qt.TextFlag.TextSingleLine, text[:25])
                y_offset += 12

        # 绘制页码
        painter.setPen(QPen(COLOR_TEXT))
        font = QFont()
        font.setPointSize(5)
        painter.setFont(font)
        page_rect = QRectF(0, self.thumb_height - 14, self.thumb_width - 4, 12)
        painter.drawText(page_rect, Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter,
                         f"{index + 1}/{total}")

        # 绘制底部装饰线
        painter.setPen(QPen(COLOR_ACCENT, 1.5))
        painter.drawLine(0, self.thumb_height - 2, self.thumb_width, self.thumb_height - 2)

        painter.end()
        return pixmap

    def _render_from_data(self, slide_data: dict, index: int, total: int) -> QPixmap:
        """从 slides_data 格式渲染缩略图"""
        pixmap = QPixmap(self.thumb_width, self.thumb_height)
        pixmap.fill(COLOR_BG)

        painter = QPainter(pixmap)
        painter.setRenderHint(QPainter.RenderHint.Antialiasing)

        title = slide_data.get("title", "")
        items = slide_data.get("items", [])

        # 绘制标题区域
        if title:
            painter.fillRect(0, 0, self.thumb_width, self.thumb_height // 3, COLOR_LIGHT_BG)
            painter.setPen(QPen(COLOR_TITLE))
            font = QFont()
            font.setPointSize(7)
            font.setBold(True)
            painter.setFont(font)
            title_rect = QRectF(4, 2, self.thumb_width - 8, self.thumb_height // 3 - 4)
            painter.drawText(title_rect, Qt.TextFlag.TextWordWrap | Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter, title[:30])

        # 绘制正文
        if items:
            painter.setPen(QPen(COLOR_TEXT))
            font = QFont()
            font.setPointSize(5)
            painter.setFont(font)
            y_offset = self.thumb_height // 3 + 4
            for item in items[:4]:
                text = item.get("text", "") if isinstance(item, dict) else str(item)
                text_rect = QRectF(6, y_offset, self.thumb_width - 12, 12)
                painter.drawText(text_rect, Qt.TextFlag.TextSingleLine, text[:25])
                y_offset += 12

        # 页码
        painter.setPen(QPen(COLOR_TEXT))
        font = QFont()
        font.setPointSize(5)
        painter.setFont(font)
        page_rect = QRectF(0, self.thumb_height - 14, self.thumb_width - 4, 12)
        painter.drawText(page_rect, Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter,
                         f"{index + 1}/{total}")

        # 底部装饰线
        painter.setPen(QPen(COLOR_ACCENT, 1.5))
        painter.drawLine(0, self.thumb_height - 2, self.thumb_width, self.thumb_height - 2)

        painter.end()
        return pixmap
