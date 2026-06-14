#!/usr/bin/env python3
"""
pptx-renderer 快速验证脚本 v2

修复项：
- 拖拽排序：rowsMoved 后不 clear 列表，只就地更新文字
- 双击编辑：增加 Edit 按钮作为可靠备选
- 自由拖动：pptx-renderer 是只读渲染器，不支持元素级拖拽
"""

import base64
import os
import sys
import threading
import time
from http.server import HTTPServer, SimpleHTTPRequestHandler

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", ".."))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from PyQt6.QtCore import Qt, QUrl, pyqtSignal, pyqtSlot, QObject, QTimer
from PyQt6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QListWidget, QListWidgetItem, QPushButton, QLabel, QSplitter,
    QInputDialog, QMessageBox, QAbstractItemView,
)
from PyQt6.QtWebEngineWidgets import QWebEngineView
from PyQt6.QtWebChannel import QWebChannel

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

VIEWER_DIR = os.path.dirname(os.path.abspath(__file__))
HTTP_PORT = 18923


class _QuietHandler(SimpleHTTPRequestHandler):
    def __init__(self, *a, **kw):
        super().__init__(*a, directory=VIEWER_DIR, **kw)
    def log_message(self, fmt, *args):
        pass


class ViewerBridge(QObject):
    viewer_ready = pyqtSignal(int)
    @pyqtSlot(int)
    def onViewerReady(self, count: int):
        print(f"[Bridge] viewer ready: {count} slides")
        self.viewer_ready.emit(count)


def make_test_slides():
    return [
        {"type": "title", "title": "pptx-renderer 验证", "subtitle": "五项核心能力测试"},
        {"type": "content", "title": "能力清单", "layout": "text_only",
         "items": [
             {"level": 0, "text": "1. 实时渲染：生成 .pptx 后热加载"},
             {"level": 0, "text": "2. 拖拽排序：左侧列表拖动调整顺序"},
             {"level": 0, "text": "3. 点击编辑：选中后点 Edit 按钮修改标题"},
             {"level": 0, "text": "4. 新增删除：按钮操作增删幻灯片"},
             {"level": 0, "text": "5. 自由编排：pptx-renderer 为只读，不支持"},
         ]},
        {"type": "content", "title": "数据页", "layout": "text_only",
         "items": [
             {"level": 0, "text": "Q1: 2.8 亿 | Q2: 3.1 亿 | Q3: 3.5 亿"},
             {"level": 1, "text": "持续增长趋势"},
         ]},
    ]


def slides_to_pptx_bytes(slides_data):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for sd in slides_data:
        stype = sd.get("type", "content")
        if stype == "title":
            sl = prs.slides.add_slide(prs.slide_layouts[0])
            tb = sl.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11.333), Inches(1.5))
            tf = tb.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]; p.text = sd.get("title", "")
            p.font.size = Pt(48); p.font.bold = True; p.font.color.rgb = RGBColor(255, 255, 255)
            p.alignment = PP_ALIGN.CENTER
            bg = sl.background.fill; bg.solid(); bg.fore_color.rgb = RGBColor(0, 82, 204)
            sub = sd.get("subtitle", "")
            if sub:
                tb2 = sl.shapes.add_textbox(Inches(1), Inches(4.2), Inches(11.333), Inches(0.8))
                tf2 = tb2.text_frame; tf2.word_wrap = True
                p2 = tf2.paragraphs[0]; p2.text = sub
                p2.font.size = Pt(24); p2.font.color.rgb = RGBColor(255, 255, 255); p2.alignment = PP_ALIGN.CENTER
        else:
            sl = prs.slides.add_slide(prs.slide_layouts[1])
            tb = sl.shapes.add_textbox(Inches(1), Inches(0.5), Inches(11.333), Inches(1.2))
            tf = tb.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]; p.text = sd.get("title", "")
            p.font.size = Pt(36); p.font.bold = True; p.font.color.rgb = RGBColor(0, 82, 204)
            items = sd.get("items", [])
            if items:
                tb2 = sl.shapes.add_textbox(Inches(1), Inches(1.8), Inches(11.333), Inches(5.0))
                tf2 = tb2.text_frame; tf2.word_wrap = True
                for j, item in enumerate(items):
                    pp = tf2.paragraphs[0] if j == 0 else tf2.add_paragraph()
                    pp.text = item.get("text", ""); pp.level = item.get("level", 0); pp.space_after = Pt(8)
                    if item.get("level", 0) == 0:
                        pp.font.size = Pt(20); pp.font.color.rgb = RGBColor(15, 25, 45)
                    else:
                        pp.font.size = Pt(16); pp.font.color.rgb = RGBColor(60, 64, 67)
    import io
    buf = io.BytesIO(); prs.save(buf); return buf.getvalue()


class TestViewerWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("pptx-renderer 验证 v2")
        self.resize(1400, 800)
        self.slides_data = make_test_slides()
        self.bridge = ViewerBridge()
        self.bridge.viewer_ready.connect(lambda c: self._set_status(f"渲染完成: {c} 页"))
        self._init_ui()
        self._refresh_list()
        QTimer.singleShot(1500, self._load_pptx)

    def _init_ui(self):
        central = QWidget(); self.setCentralWidget(central)
        main_layout = QHBoxLayout(central); main_layout.setContentsMargins(0, 0, 0, 0)
        splitter = QSplitter(Qt.Orientation.Horizontal)

        # ── 左面板 ──
        left = QWidget()
        left_layout = QVBoxLayout(left); left_layout.setContentsMargins(4, 4, 4, 4)

        # 按钮行
        btn_row = QHBoxLayout()
        self.btn_add = QPushButton("+ 新增"); self.btn_add.clicked.connect(self._add_slide)
        self.btn_del = QPushButton("- 删除"); self.btn_del.clicked.connect(self._del_slide)
        self.btn_edit = QPushButton("✏ 编辑"); self.btn_edit.clicked.connect(self._edit_selected)
        self.btn_reload = QPushButton("🔄 渲染"); self.btn_reload.clicked.connect(self._load_pptx)
        for b in [self.btn_add, self.btn_del, self.btn_edit, self.btn_reload]:
            btn_row.addWidget(b)
        left_layout.addLayout(btn_row)

        # 幻灯片列表
        self.slide_list = QListWidget()
        self.slide_list.setDragDropMode(QAbstractItemView.DragDropMode.InternalMove)
        self.slide_list.setDefaultDropAction(Qt.DropAction.MoveAction)
        self.slide_list.itemDoubleClicked.connect(self._on_double_click)
        self.slide_list.model().rowsMoved.connect(self._on_rows_moved)
        left_layout.addWidget(self.slide_list, stretch=1)

        # 状态
        self.status_label = QLabel("就绪")
        self.status_label.setStyleSheet("color: #888; font-size: 11px;")
        left_layout.addWidget(self.status_label)

        # 能力说明
        hint = QLabel("拖拽列表条目可调整顺序\n选中条目后点 Edit 或双击可编辑")
        hint.setStyleSheet("color: #666; font-size: 10px;")
        hint.setWordWrap(True)
        left_layout.addWidget(hint)

        splitter.addWidget(left)

        # ── 右面板：QWebEngineView ──
        self.web_view = QWebEngineView()
        splitter.addWidget(self.web_view)
        splitter.setSizes([260, 1140])
        main_layout.addWidget(splitter)

        # QWebChannel
        page = self.web_view.page()
        self.channel = QWebChannel(page)
        self.channel.registerObject("pyBridge", self.bridge)
        page.setWebChannel(self.channel)
        page.setUrl(QUrl(f"http://127.0.0.1:{HTTP_PORT}/test_viewer.html"))

    # ── 列表管理 ──

    def _refresh_list(self):
        """完整刷新列表"""
        self.slide_list.blockSignals(True)
        self.slide_list.clear()
        for i, sd in enumerate(self.slides_data):
            self._add_list_item(i, sd)
        self.slide_list.blockSignals(False)

    def _add_list_item(self, idx, sd):
        """添加一个列表条目"""
        title = sd.get("title", f"第 {idx+1} 页")
        stype = sd.get("type", "content")
        icon = "🎯" if stype == "title" else "📄"
        item = QListWidgetItem(f"{icon} {idx+1}. {title}")
        item.setFlags(
            Qt.ItemFlag.ItemIsSelectable | Qt.ItemFlag.ItemIsEnabled
            | Qt.ItemFlag.ItemIsDragEnabled | Qt.ItemFlag.ItemIsDropEnabled
        )
        item.setData(Qt.ItemDataRole.UserRole, idx)
        self.slide_list.addItem(item)

    def _update_list_item_text(self, item, idx, sd):
        """就地更新条目文字（不触发 clear）"""
        title = sd.get("title", f"第 {idx+1} 页")
        stype = sd.get("type", "content")
        icon = "🎯" if stype == "title" else "📄"
        item.setText(f"{icon} {idx+1}. {title}")
        item.setData(Qt.ItemDataRole.UserRole, idx)

    def _renumber_items(self):
        """重新编号所有条目（拖拽后调用）"""
        for i in range(self.slide_list.count()):
            item = self.slide_list.item(i)
            sd_idx = item.data(Qt.ItemDataRole.UserRole)
            if sd_idx is not None and 0 <= sd_idx < len(self.slides_data):
                sd = self.slides_data[sd_idx]
                self._update_list_item_text(item, i, sd)

    # ── PPTX 生成与加载 ──

    def _load_pptx(self):
        t0 = time.time()
        pptx_bytes = slides_to_pptx_bytes(self.slides_data)
        b64 = base64.b64encode(pptx_bytes).decode("ascii")
        elapsed = (time.time() - t0) * 1000
        print(f"[DEBUG] _load_pptx: {len(self.slides_data)} slides, {len(pptx_bytes)}B, {elapsed:.0f}ms")

        js = f"""
        (async function() {{
            const b64 = "{b64}";
            const bin = atob(b64);
            const bytes = new Uint8Array(bin.length);
            for (let i = 0; i < bin.length; i++) bytes[i] = bin.charCodeAt(i);
            await loadPptx(bytes.buffer);
            return "ok";
        }})();
        """
        def _cb(result):
            print(f"[DEBUG] JS callback: {result}")
        self.web_view.page().runJavaScript(js, 0, _cb)
        self._set_status(f"生成 .pptx: {elapsed:.0f}ms | {len(pptx_bytes)}B")

    # ── 交互处理 ──

    def _edit_selected(self):
        """编辑选中幻灯片的标题"""
        item = self.slide_list.currentItem()
        if not item:
            self._set_status("请先选中一页幻灯片")
            return
        idx = self.slide_list.row(item)
        if idx < 0 or idx >= len(self.slides_data):
            return
        print(f"[DEBUG] edit: row={idx}, title='{self.slides_data[idx].get('title','')}'")
        old_title = self.slides_data[idx].get("title", "")
        new_title, ok = QInputDialog.getText(self, "编辑标题", "新标题:", text=old_title)
        print(f"[DEBUG] edit result: ok={ok}, new_title='{new_title}'")
        if ok and new_title != old_title:
            self.slides_data[idx]["title"] = new_title
            # 就地更新列表文字
            self._update_list_item_text(item, idx, self.slides_data[idx])
            self._load_pptx()

    def _on_double_click(self, item):
        """双击 = 编辑"""
        print(f"[DEBUG] double click: row={self.slide_list.row(item)}")
        self._edit_selected()

    def _on_rows_moved(self, parent, start, end, dest, row):
        """拖拽排序完成"""
        print(f"[DEBUG] rowsMoved: start={start} end={end} dest={dest} row={row}")
        if start == row:
            return

        # 从列表控件读取新顺序
        new_order = []
        for i in range(self.slide_list.count()):
            item = self.slide_list.item(i)
            orig = item.data(Qt.ItemDataRole.UserRole)
            if orig is not None:
                new_order.append(orig)
        print(f"[DEBUG] new order: {new_order}")

        if len(new_order) != len(self.slides_data):
            print(f"[DEBUG] order length mismatch, skipping")
            return

        # 重排 slides_data
        self.slides_data = [self.slides_data[i] for i in new_order]
        # 就地更新列表编号（不清空列表！）
        self._renumber_items()
        # 重新渲染预览
        self._load_pptx()
        self._set_status(f"拖拽排序完成: {new_order}")

    def _add_slide(self):
        print("[DEBUG] add slide")
        n = len(self.slides_data) + 1
        self.slides_data.append({
            "type": "content", "title": f"新页面 {n}", "layout": "text_only",
            "items": [{"level": 0, "text": "新内容"}],
        })
        # 就地添加列表条目
        self._add_list_item(n - 1, self.slides_data[-1])
        self._load_pptx()

    def _del_slide(self):
        print("[DEBUG] del slide")
        idx = self.slide_list.currentRow()
        if idx < 0 or idx >= len(self.slides_data):
            self._set_status("请先选中一页")
            return
        if len(self.slides_data) <= 1:
            self._set_status("至少保留 1 页")
            return
        self.slides_data.pop(idx)
        # 就地删除列表条目
        self.slide_list.takeItem(idx)
        self._renumber_items()
        self._load_pptx()

    def _set_status(self, msg):
        self.status_label.setText(msg)
        print(f"[STATUS] {msg}")


if __name__ == "__main__":
    # HTTP 服务
    server = HTTPServer(("127.0.0.1", HTTP_PORT), _QuietHandler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    print(f"[HTTP] {VIEWER_DIR} on port {HTTP_PORT}")

    app = QApplication(sys.argv)
    win = TestViewerWindow()
    win.show()
    sys.exit(app.exec())
