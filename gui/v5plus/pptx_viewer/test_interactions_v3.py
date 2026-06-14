#!/usr/bin/env python3
"""pptx-renderer 交互能力验证 v3 - overlay 方案"""
import base64, os, sys, threading, time
from http.server import HTTPServer, SimpleHTTPRequestHandler

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", ".."))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from PyQt6.QtCore import Qt, QUrl, pyqtSignal, pyqtSlot, QObject, QTimer
from PyQt6.QtWidgets import QApplication, QMainWindow, QVBoxLayout, QWidget, QLabel
from PyQt6.QtWebEngineWidgets import QWebEngineView
from PyQt6.QtWebChannel import QWebChannel

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

VIEWER_DIR = os.path.dirname(os.path.abspath(__file__))
HTTP_PORT = 18924

class _H(SimpleHTTPRequestHandler):
    def __init__(self, *a, **kw): super().__init__(*a, directory=VIEWER_DIR, **kw)
    def log_message(self, *a): pass

class Bridge(QObject):
    viewer_ready = pyqtSignal(int)
    @pyqtSlot(int)
    def onViewerReady(self, c): self.viewer_ready.emit(c)

def make_test_pptx():
    """3 页测试 PPT：封面 + 多文本元素 + 表格"""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 封面
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    tb = s1.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11.333), Inches(1.5))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = "交互验证测试"; p.font.size = Pt(48); p.font.bold = True
    p.font.color.rgb = RGBColor(255,255,255); p.alignment = PP_ALIGN.CENTER
    bg = s1.background.fill; bg.solid(); bg.fore_color.rgb = RGBColor(0,82,204)

    # 内容页 - 多个独立文本框
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    tb = s2.shapes.add_textbox(Inches(1), Inches(0.5), Inches(11), Inches(1))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = "多元素交互测试"; p.font.size = Pt(36); p.font.bold = True; p.font.color.rgb = RGBColor(0,82,204)

    # 独立文本框 1
    tb1 = s2.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(1))
    tf1 = tb1.text_frame; tf1.word_wrap = True
    p1 = tf1.paragraphs[0]; p1.text = "可拖拽的文本框 A"; p1.font.size = Pt(20); p1.font.color.rgb = RGBColor(15,25,45)

    # 独立文本框 2
    tb2 = s2.shapes.add_textbox(Inches(7), Inches(2), Inches(5), Inches(1))
    tf2 = tb2.text_frame; tf2.word_wrap = True
    p2 = tf2.paragraphs[0]; p2.text = "可编辑的文本框 B"; p2.font.size = Pt(20); p2.font.color.rgb = RGBColor(15,25,45)

    # 独立文本框 3
    tb3 = s2.shapes.add_textbox(Inches(1), Inches(4), Inches(11), Inches(2))
    tf3 = tb3.text_frame; tf3.word_wrap = True
    p3 = tf3.paragraphs[0]; p3.text = "底部大文本区域 - 试试双击编辑和拖拽移动"; p3.font.size = Pt(18)
    p3.font.color.rgb = RGBColor(60,64,67)

    # 表格页
    s3 = prs.slides.add_slide(prs.slide_layouts[1])
    tb = s3.shapes.add_textbox(Inches(1), Inches(0.5), Inches(11), Inches(1))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = "表格测试"; p.font.size = Pt(36); p.font.bold = True; p.font.color.rgb = RGBColor(0,82,204)

    from pptx.util import Emu
    table = s3.shapes.add_table(4, 3, Inches(1), Inches(2), Inches(11), Inches(4)).table
    table.columns[0].width = Inches(3)
    table.columns[1].width = Inches(4)
    table.columns[2].width = Inches(4)
    headers = ["指标", "Q1", "Q2"]
    data = [["营收", "2.8亿", "3.1亿"], ["利润", "0.5亿", "0.7亿"], ["增长率", "15%", "21%"]]
    for i, h in enumerate(headers):
        cell = table.cell(0, i); cell.text = h
        for p in cell.text_frame.paragraphs:
            p.font.bold = True; p.font.size = Pt(16)
    for r, row_data in enumerate(data):
        for c, val in enumerate(row_data):
            cell = table.cell(r+1, c); cell.text = val
            for p in cell.text_frame.paragraphs: p.font.size = Pt(14)

    import io
    buf = io.BytesIO(); prs.save(buf); return buf.getvalue()


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("pptx-renderer 交互验证 v3 - overlay")
        self.resize(1400, 800)
        self.bridge = Bridge()
        central = QWidget(); self.setCentralWidget(central)
        layout = QVBoxLayout(central); layout.setContentsMargins(0,0,0,0)
        self.web = QWebEngineView(); layout.addWidget(self.web)
        page = self.web.page()
        ch = QWebChannel(page); ch.registerObject("pyBridge", self.bridge)
        page.setWebChannel(ch)
        page.setUrl(QUrl(f"http://127.0.0.1:{HTTP_PORT}/test_interactions_v3.html"))
        QTimer.singleShot(2000, self._load)

    def _load(self):
        pptx_bytes = make_test_pptx()
        b64 = base64.b64encode(pptx_bytes).decode("ascii")
        print(f"[DEBUG] loading {len(pptx_bytes)} bytes")
        js = f"""
        (async function() {{
            const b64 = "{b64}";
            const bin = atob(b64);
            const bytes = new Uint8Array(bin.length);
            for (let i = 0; i < bin.length; i++) bytes[i] = bin.charCodeAt(i);
            await loadPptx(bytes.buffer);
            return "ok";
        }})();
        """
        self.web.page().runJavaScript(js, 0, lambda r: print(f"[DEBUG] JS load: {r}"))
        # 5秒后读取日志
        QTimer.singleShot(5000, self._read_log)

    def _read_log(self):
        js = "document.getElementById('log')?.textContent || 'no log'"
        self.web.page().runJavaScript(js, 0, lambda r: print(f"[LOG DUMP]\n{r}"))


if __name__ == "__main__":
    server = HTTPServer(("127.0.0.1", HTTP_PORT), _H)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    print(f"[HTTP] port {HTTP_PORT}")
    app = QApplication(sys.argv)
    win = MainWindow(); win.show()
    print("窗口已打开，请在右侧测试：")
    print("  1. 鼠标悬停 → overlay 蓝色边框高亮（不会被遮挡）")
    print("  2. 点击元素 → overlay 蓝色实线框 + 阴影")
    print("  3. 切换到✋拖拽模式 → 拖拽 overlay 框，永远不会被遮挡")
    print("  4. 切换到✏编辑模式 → 双击文字编辑，字体继承原始")
    print("  5. 点击🔍检查DOM → 查看DOM结构")
    sys.exit(app.exec())
